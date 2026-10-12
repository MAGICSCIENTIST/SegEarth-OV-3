"""Shared metric conversion and Excel reporting utilities.

The helpers in this module are intentionally independent of MMSeg so they can
be reused by evaluation and the lightweight parameter-training entry points.
"""

from __future__ import annotations

import math
import os
import os.path as osp
import sys
import time
import traceback
from typing import Mapping, Sequence

import openpyxl


SUMMARY_HEADERS = [
    'Model', 'Dataset', 'aAcc', 'mIoU', 'mAcc', 'mFscore',
    'mPrecision', 'mRecall',
    'InitialaAcc', 'aAccGain',
    'InitialmIoU', 'mIoUGain',
    'InitialmAcc', 'mAccGain',
    'InitialmFscore', 'mFscoreGain',
    'InitialmPrecision', 'mPrecisionGain',
    'InitialmRecall', 'mRecallGain',
    'Epoch', 'BestEpoch', 'Split', 'Loss',
]
CLASS_HEADERS = [
    'Model', 'Dataset', 'Class', 'IoU', 'Acc', 'Fscore', 'Precision',
    'Recall', 'TP', 'FP', 'FN', 'PredPixels', 'GTPixels', 'PredGTAreaRatio',
    'InitialIoU', 'IoUGain', 'InitialAcc', 'AccGain',
    'InitialFscore', 'FscoreGain', 'InitialPrecision', 'PrecisionGain',
    'InitialRecall', 'RecallGain', 'InitialTP',
    'MaxBaselineIoU', 'IoUDeltaVsMax',
    'Epoch', 'BestEpoch', 'Split', 'Threshold',
]
PROMPT_HEADERS = [
    'Model', 'Dataset', 'Prompt', 'ConceptID', 'ConceptName', 'IoU', 'Acc',
    'Fscore', 'Precision', 'Recall', 'TP', 'FP', 'FN', 'PredPixels',
    'GTPixels', 'PredGTAreaRatio',
    'InitialIoU', 'IoUGain', 'InitialAcc', 'AccGain',
    'InitialFscore', 'FscoreGain', 'InitialPrecision', 'PrecisionGain',
    'InitialRecall', 'RecallGain', 'InitialTP',
    'Epoch', 'BestEpoch', 'Split', 'Threshold',
]


def safe_ratio(numerator, denominator):
    denominator = float(denominator)
    return float(numerator) / denominator if denominator else float('nan')


def counts_to_metrics(intersect, union, predicted, target, percent=True):
    """Convert aggregate binary pixel counts into a serializable metric row."""
    tp = float(intersect)
    pred = float(predicted)
    gt = float(target)
    fp = pred - tp
    fn = gt - tp
    scale = 100.0 if percent else 1.0
    precision = safe_ratio(tp, pred)
    recall = safe_ratio(tp, gt)
    fscore = safe_ratio(2.0 * precision * recall, precision + recall)
    return dict(
        IoU=safe_ratio(tp, union) * scale,
        Acc=recall * scale,
        Fscore=fscore * scale,
        Precision=precision * scale,
        Recall=recall * scale,
        TP=int(round(tp)),
        FP=int(round(fp)),
        FN=int(round(fn)),
        PredPixels=int(round(pred)),
        GTPixels=int(round(gt)),
        PredGTAreaRatio=safe_ratio(pred, gt),
    )


def summarize_metric_rows(rows: Sequence[Mapping], epoch=None,
                          split=None, loss=None, best_epoch=None):
    """Build eval-style macro and overall metrics from per-class rows.

    Per-class percentage fields are averaged with undefined values ignored.
    ``aAcc`` is the micro pixel accuracy/recall derived from total TP and total
    GT pixels, matching the class-count convention used by the trainers.
    """
    rows = list(rows or [])

    def mean_defined(key):
        values = []
        for row in rows:
            value = row.get(key)
            if value is None:
                continue
            value = float(value)
            if math.isfinite(value):
                values.append(value)
        return sum(values) / len(values) if values else None

    total_tp = sum(float(row.get('TP') or 0) for row in rows)
    has_initial_tp = any(row.get('InitialTP') is not None for row in rows)
    total_initial_tp = sum(float(row.get('InitialTP') or 0) for row in rows)
    total_gt = sum(float(row.get('GTPixels') or 0) for row in rows)
    metrics = ('IoU', 'Acc', 'Fscore', 'Precision', 'Recall')
    final_means = {key: mean_defined(key) for key in metrics}
    initial_means = {
        key: mean_defined(f'Initial{key}') for key in metrics
    }
    gains = {
        key: (None if final_means[key] is None or initial_means[key] is None
              else final_means[key] - initial_means[key])
        for key in metrics
    }
    aacc = 100.0 * total_tp / total_gt if total_gt else None
    initial_aacc = (100.0 * total_initial_tp / total_gt
                    if total_gt and has_initial_tp else None)
    return dict(
        aAcc=aacc,
        mIoU=final_means['IoU'],
        mAcc=final_means['Acc'],
        mFscore=final_means['Fscore'],
        mPrecision=final_means['Precision'],
        mRecall=final_means['Recall'],
        InitialaAcc=initial_aacc,
        aAccGain=(None if aacc is None or initial_aacc is None
                  else aacc - initial_aacc),
        InitialmIoU=initial_means['IoU'],
        mIoUGain=gains['IoU'],
        InitialmAcc=initial_means['Acc'],
        mAccGain=gains['Acc'],
        InitialmFscore=initial_means['Fscore'],
        mFscoreGain=gains['Fscore'],
        InitialmPrecision=initial_means['Precision'],
        mPrecisionGain=gains['Precision'],
        InitialmRecall=initial_means['Recall'],
        mRecallGain=gains['Recall'],
        Epoch=epoch,
        BestEpoch=best_epoch,
        Split=split,
        Loss=loss,
    )


def format_duration(seconds):
    seconds = max(float(seconds), 0.0)
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    if hours > 0:
        return f'{hours:02d}:{minutes:02d}:{secs:02d}'
    return f'{minutes:02d}:{secs:02d}'


class ProgressReporter:
    """Reusable elapsed/ETA/expected-finish estimator for training logs."""

    def __init__(self, total=None, started_at=None):
        self.total = None if total is None else int(total)
        self.started_at = time.time() if started_at is None else float(started_at)

    def snapshot(self, completed, total=None, now=None):
        total = self.total if total is None else int(total)
        completed = int(completed)
        now = time.time() if now is None else float(now)
        elapsed = max(now - self.started_at, 0.0)
        if total is None or total <= 0 or completed <= 0:
            eta = None
            expected_finish = None
        else:
            remaining = max(total - completed, 0)
            eta = elapsed / completed * remaining
            expected_finish = now + eta
        return dict(
            completed=completed,
            total=total,
            elapsed_seconds=elapsed,
            eta_seconds=eta,
            expected_finish_timestamp=expected_finish)

    def format(self, completed, total=None, now=None):
        state = self.snapshot(completed, total=total, now=now)
        eta = state['eta_seconds']
        expected = state['expected_finish_timestamp']
        return (
            f'elapsed={format_duration(state["elapsed_seconds"])} '
            f'eta={"pending" if eta is None else format_duration(eta)} '
            f'expected_finish={"pending" if expected is None else time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(expected))}')


class _TeeStream:
    def __init__(self, stream, *log_files):
        self.stream = stream
        self.log_files = log_files

    def write(self, text):
        self.stream.write(text)
        for log_file in self.log_files:
            log_file.write(text)
        return len(text)

    def flush(self):
        self.stream.flush()
        for log_file in self.log_files:
            log_file.flush()

    def __getattr__(self, name):
        return getattr(self.stream, name)


class TrainingOutputReporter:
    """Mirror one training run to terminal, log.txt, and run.log."""

    def __init__(self, output_dir, run_name):
        self.output_dir = osp.abspath(output_dir)
        self.run_name = str(run_name)
        self.timer = ProgressReporter()
        self._log_files = []
        self._stdout = None
        self._stderr = None

    def __enter__(self):
        os.makedirs(self.output_dir, exist_ok=True)
        self._log_files = [
            open(osp.join(self.output_dir, name), 'a', encoding='utf-8',
                 buffering=1)
            for name in ('log.txt', 'run.log')
        ]
        self._stdout, self._stderr = sys.stdout, sys.stderr
        sys.stdout = _TeeStream(self._stdout, *self._log_files)
        sys.stderr = _TeeStream(self._stderr, *self._log_files)
        print(f'\n=== {self.run_name} started '
              f'{time.strftime("%Y-%m-%d %H:%M:%S")} ===', flush=True)
        print(
            f'log_file={osp.join(self.output_dir, "log.txt")} '
            f'run_log={osp.join(self.output_dir, "run.log")} '
            'expected_finish=pending', flush=True)
        return self

    def __exit__(self, exc_type, exc_value, exc_traceback):
        try:
            if exc_type is not None:
                print(f'=== {self.run_name} failed ===', file=sys.stderr,
                      flush=True)
                traceback.print_exception(
                    exc_type, exc_value, exc_traceback, file=sys.stderr)
            print(
                f'=== {self.run_name} ended '
                f'{time.strftime("%Y-%m-%d %H:%M:%S")} '
                f'{self.timer.format(1, total=1)} ===', flush=True)
        finally:
            sys.stdout, sys.stderr = self._stdout, self._stderr
            for log_file in self._log_files:
                log_file.close()
        return False


def _clean_value(value):
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def _ensure_sheet(workbook, name, headers):
    if name == 'summary':
        sheet = workbook[name] if name in workbook.sheetnames else workbook.active
        sheet.title = name
    else:
        sheet = (workbook[name] if name in workbook.sheetnames
                 else workbook.create_sheet(name))
    # Preserve existing files and columns; append any newly introduced fields.
    existing = [cell.value for cell in sheet[1]] if sheet.max_row else []
    if not any(existing):
        existing = []
    for header in headers:
        if header not in existing:
            existing.append(header)
    for column, header in enumerate(existing, 1):
        sheet.cell(1, column, header)
    return sheet, existing


def _append_rows(sheet, headers, rows, common):
    for row in rows or []:
        values = dict(common)
        values.update(row)
        sheet.append([_clean_value(values.get(header)) for header in headers])


def write_metrics_workbook(file_path: str,
                           summaries: Sequence[Mapping] | None = None,
                           class_metrics: Sequence[Mapping] | None = None,
                           prompt_metrics: Sequence[Mapping] | None = None,
                           model: str | None = None,
                           dataset: str | None = None):
    """Append metric rows while retaining legacy sheets and columns."""
    os.makedirs(osp.dirname(osp.abspath(file_path)), exist_ok=True)
    try:
        workbook = openpyxl.load_workbook(file_path)
    except FileNotFoundError:
        workbook = openpyxl.Workbook()
    common = {'Model': model, 'Dataset': dataset}
    if summaries:
        sheet, headers = _ensure_sheet(workbook, 'summary', SUMMARY_HEADERS)
        _append_rows(sheet, headers, summaries, common)
    if class_metrics:
        sheet, headers = _ensure_sheet(workbook, 'per_class', CLASS_HEADERS)
        _append_rows(sheet, headers, class_metrics, common)
    if prompt_metrics:
        sheet, headers = _ensure_sheet(workbook, 'per_prompt', PROMPT_HEADERS)
        _append_rows(sheet, headers, prompt_metrics, common)
    workbook.save(file_path)


class MetricsReporter:
    """Shared facade used by evaluation and parameter-training scripts."""

    counts_to_metrics = staticmethod(counts_to_metrics)
    format_duration = staticmethod(format_duration)
    summarize_rows = staticmethod(summarize_metric_rows)
    write_workbook = staticmethod(write_metrics_workbook)
