"""Batch final multiclass evaluation of completed CX-B runs (no training).

python eval_batch_cx_b.py --root E:/output/SEGOV3/weights/stage_cx_b
Results: <run>/eval_multiclass/val/{metrics.json,results.xlsx,summary.csv,...}
Two decision rules share one forward pass: saved eval_cx_b defaults and a
common threshold/background rule (0.5/0, matching Eval LoveDA by default).
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
import gc
import hashlib
import json
import os
from pathlib import Path
import time
import traceback

import numpy as np
from PIL import Image
import torch
from openpyxl import Workbook

from cx_b.checkpoints import atomic_json_save
from cx_b.configuration import refiner_settings
from cx_b.refiners import JOINT_FORMATS, evaluation_refiner_options, ablation_folder
from cx_b.reporting import prediction_changes
from eval_cx_b import (
    NODATA, REPO_ROOT, build_head_source, classify_scores, collect_samples,
    load_run, map_gt, metric_rows, predict_image, resolve_path,
)

VERSION = 2
LOVEDA_CLASSES = ['background', 'building', 'road', 'water', 'barren', 'forest', 'agricultural']


def boolean_value(value):
    if value.lower() in ('true', '1', 'yes'):
        return True
    if value.lower() in ('false', '0', 'no'):
        return False
    raise argparse.ArgumentTypeError('Expected true or false')


def parse_args(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root', type=Path, required=True, help='Stage, combination, or seed directory.')
    p.add_argument('--out-root', type=Path, help='Optional separate output tree, e.g. for smoke checks.')
    p.add_argument('--split', choices=['val', 'train', 'test'], default='val')
    p.add_argument('--device', default='cuda' if torch.cuda.is_available() else 'cpu')
    p.add_argument('--common-threshold', type=float, default=.5)
    p.add_argument('--common-background', type=int, default=0)
    p.add_argument('--base-dir', type=Path, help='Optional cfg_loveda output root; read its result/*.png.')
    p.add_argument('--use-best', type=boolean_value, nargs='?', const=True, default=True,
                   help='true (default): original best.pth assembly; false: evaluate every module pth and the winners assembly.')
    p.add_argument('--no-use-best', dest='use_best', action='store_false')
    p.add_argument('--data-root')
    p.add_argument('--image-dir')
    p.add_argument('--gt-dir')
    p.add_argument('--prompt-head-cache', help='Override raw cache root for every selected cached run.')
    p.add_argument('--include', nargs='+', help='Select combination directory names (or seed paths relative to root).')
    p.add_argument('--max-images', type=int, default=-1, help='-1: all images; positive values create a separate smoke folder.')
    p.add_argument('--progress-every', type=int, default=20)
    p.add_argument('--save-every', type=int, default=20, help='Save small metric accumulators for image-level resume.')
    p.add_argument('--cpu-threads', type=int, default=4)
    p.add_argument('--force', action='store_true', help='Recompute instead of skipping or resuming matching results.')
    p.add_argument('--joint-refiner', choices=('none', 'ov_class_spatial'), default=None)
    p.add_argument('--joint-refiner-strength', type=float, default=None)
    p.add_argument('--list-only', action='store_true', help='List runs/status without inference or writing files.')
    args = p.parse_args(argv)
    if not 0 <= args.common_threshold <= 1:
        p.error('--common-threshold must be in [0,1]')
    if args.max_images == 0 or args.max_images < -1:
        p.error('--max-images must be positive or -1')
    if min(args.progress_every, args.save_every, args.cpu_threads) < 1:
        p.error('Progress, save intervals and CPU threads must be positive')
    args.root = args.root.expanduser().resolve()
    if not args.root.is_dir():
        p.error(f'Root directory does not exist: {args.root}')
    if args.out_root:
        args.out_root = args.out_root.expanduser().resolve()
    if args.base_dir and not (args.base_dir/'result').is_dir():
        p.error('--base-dir must contain the historical result directory')
    return args


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def discover_runs(args):
    entries = []
    for snapshot_path in sorted(args.root.rglob('run_config.json')):
        directory = snapshot_path.parent
        if args.include and directory.parent.name not in args.include and directory.relative_to(args.root).as_posix() not in args.include:
            continue
        row = dict(directory=str(directory), combination=directory.parent.name, seed=directory.name)
        try:
            run = read_json(snapshot_path)
            cx = run['config']['cx_b']
            row.update({k: cx[k] for k in ('prompt_method', 'concept_method', 'loss_method', 'data_mode')})
            marker = directory/'completed.json'
            if not marker.exists():
                row['status'] = 'skipped_incomplete'
            else:
                done = read_json(marker)
                if done.get('status') != 'complete':
                    row['status'] = 'skipped_incomplete'
                elif not run.get('fingerprint') or done.get('fingerprint') != run['fingerprint']:
                    raise ValueError('completed.json and run_config.json fingerprints differ')
                else:
                    row['status'] = 'ready'
        except Exception as error:
            row.update(status='failed', error=f'{type(error).__name__}: {error}')
        entries.append(row)
    if not entries:
        raise ValueError('No matching CX-B run_config.json found')
    return entries


def output_directory(directory, args):
    base = args.out_root/directory.relative_to(args.root) if args.out_root else directory
    split = args.split if args.max_images == -1 else f'{args.split}_first_{args.max_images}'
    path = base/'eval_multiclass'/split
    return path/ablation_folder(args) if ablation_folder(args) else path


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024*1024), b''):
            digest.update(block)
    return digest.hexdigest()


def stat_signature(paths):
    digest = hashlib.sha256()
    for path in paths:
        path = Path(path).resolve()
        stat = path.stat()
        digest.update(json.dumps([str(path), stat.st_size, stat.st_mtime_ns]).encode())
    return digest.hexdigest()


def evaluation_signature(directory, run, samples, args, policies):
    cx = run['config']['cx_b']
    folder = 'dynamic_prompt' if cx['prompt_method'] == 'dynamic_head' else 'prompt_checkpoints'
    checkpoints = sorted((directory/folder).glob('*/best.pth')) + sorted((directory/'concept').glob('*/best.pth'))
    if refiner_settings(cx)[0] != 'none':
        checkpoints.append(directory/'joint/best.pth')
    data = dict(version=VERSION, snapshot_sha256=sha256(directory/'run_config.json'),
                checkpoint_sha256={p.relative_to(directory).as_posix(): sha256(p) for p in checkpoints},
                run_fingerprint=run['fingerprint'], split=args.split, max_images=args.max_images,
                device=str(args.device), torch_version=str(torch.__version__), policies=policies,
                dataset_signature=stat_signature(p for sample in samples for p in sample[:2]),
                sources={name: sha256(REPO_ROOT/name) for name in (
                    'eval_batch_cx_b.py', 'eval_cx_b.py', 'cx_b/modules.py', 'cx_b/heads.py')})
    if refiner_settings(cx)[0] != 'none' or ablation_folder(args):
        data['refiner_override'] = evaluation_refiner_options(args)
        data['sources']['cx_b/refiners.py'] = sha256(REPO_ROOT/'cx_b/refiners.py')
    if cx['data_mode'] == 'cache':
        cache_root = resolve_path(args.prompt_head_cache or run.get('inputs', {}).get('prompt_head_cache') or cx['prompt_head_cache'])
        files = sorted((cache_root/args.split).glob('prompt_*/*.pth'))
        if not files:
            raise FileNotFoundError(f'No cached heads: {cache_root/args.split}')
        data.update(cache_root=str(cache_root), cache_signature=stat_signature(files))
    else:
        concept_file = resolve_path(run['config']['model']['classname_path'])
        data['sam_inputs'] = stat_signature([REPO_ROOT/'weights/sam3/sam3.pt', concept_file])
        data['sources']['segearthov3_learnable_segmentor.py'] = sha256(
            REPO_ROOT/'segearthov3_learnable_segmentor.py')
    if args.base_dir:
        paths = [args.base_dir/'result'/p.name for p, _, _ in samples]
        if len(set(paths)) != len(paths):
            raise ValueError('Duplicate image basenames cannot match flat LoveDA baseline results')
        data['baseline_signature'] = stat_signature(paths)
    fingerprint = hashlib.sha256(json.dumps(data, sort_keys=True).encode()).hexdigest()
    return fingerprint, data


def candidate_signature(inputs, directory, overrides):
    """Bind results to the exact complete assembly, including replacement weights."""
    data = dict(inputs)
    data['checkpoint_sha256'] = dict(inputs['checkpoint_sha256'])
    for slot, path in overrides.items():
        del data['checkpoint_sha256'][slot]
        data['checkpoint_sha256'][path] = sha256(directory/path)
    data['checkpoint_overrides'] = overrides
    return hashlib.sha256(json.dumps(data, sort_keys=True).encode()).hexdigest(), data


def add_confusion(confusion, prediction, target, valid):
    count = confusion.shape[0]
    predicted = prediction[valid].astype(np.int64)
    if np.any((predicted != NODATA) & ((predicted < 0) | (predicted >= count))):
        raise ValueError('Prediction contains invalid class indices')
    predicted[predicted == NODATA] = count
    truth = target[valid].astype(np.int64)
    confusion += np.bincount(truth*(count+1)+predicted,
                             minlength=count*(count+1)).reshape(count, count+1)


def add_binary(counts, row, prediction, target, valid):
    predicted = prediction & valid
    truth = target & valid
    counts[row] += (int(np.count_nonzero(predicted & truth)),
                    int(np.count_nonzero(predicted & ~truth)),
                    int(np.count_nonzero(~predicted & truth)))


def summarize_counts(counts, names):
    rows = metric_rows(counts, names)
    for row in rows:
        for metric in ('IoU', 'Precision', 'Recall', 'Fscore'):
            row[metric] = None if row[metric] is None else row[metric]*100
        row['Acc'] = row['Recall']
        row['GTPixels'] = row['TP']+row['FN']
        row['PredPixels'] = row['TP']+row['FP']
        row['PredGTAreaRatio'] = row['PredPixels']/row['GTPixels'] if row['GTPixels'] else None
    summary = {}
    for metric in ('IoU', 'Acc', 'Precision', 'Recall', 'Fscore'):
        values = [r[metric] for r in rows if r[metric] is not None]
        summary['m'+metric] = float(np.mean(values)) if values else None
    return dict(summary=summary, per_class=rows)


def summarize_confusion(confusion, names):
    tp = np.diag(confusion)
    gt = confusion.sum(1)
    pred = confusion[:, :len(names)].sum(0)
    result = summarize_counts(np.stack([tp, pred-tp, gt-tp], axis=1), names)
    # Match MMSeg IoUMetric's F-score convention, including zero-overlap
    # classes: undefined precision/recall or 0/0 gives NaN (JSON null), and
    # nanmean excludes that class. Binary reports retain the TP/FP/FN formula.
    for row in result['per_class']:
        precision, recall = row['Precision'], row['Recall']
        row['Fscore'] = (None if precision is None or recall is None or precision+recall == 0
                         else 2*precision*recall/(precision+recall))
    fscores = [row['Fscore'] for row in result['per_class'] if row['Fscore'] is not None]
    result['summary']['mFscore'] = float(np.mean(fscores)) if fscores else None
    valid = int(confusion.sum())
    result['summary']['aAcc'] = float(tp.sum()/valid*100) if valid else None
    result.update(valid_pixels=valid, unassigned_pixels=int(confusion[:, -1].sum()),
                  confusion_matrix=confusion.tolist(),
                  confusion_columns=names+['unassigned'], confusion_rows=names)
    return result


def export_tables(output, result):
    summaries, classes, prompts = [], [], []
    for policy, metrics in result['metrics'].items():
        summaries.append(dict(policy=policy, scale='percent', images=result['images'],
                              evaluation_scope=result['evaluation_scope'], **metrics['summary']))
        destination = prompts if policy.startswith('prompt_') else classes
        destination.extend(dict(policy=policy, **row) for row in metrics['per_class'])
    book = Workbook()
    book.remove(book.active)
    for name, rows in [('summary', summaries), ('per_class', classes), ('per_prompt', prompts)]:
        headers = list(dict.fromkeys(key for row in rows for key in row))
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(k) for k in headers])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        temp = output/(name+'.csv.tmp')
        with temp.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)
        os.replace(temp, output/(name+'.csv'))
    temporary = output/'results.tmp.xlsx'
    book.save(temporary)
    book.close()
    os.replace(temporary, output/'results.xlsx')


class CandidateEvaluation:
    """Small, resumable counters for one complete prompt/concept assembly."""

    def __init__(self, output, fingerprint, inputs, count, records, samples, args, policies):
        self.output, self.fingerprint, self.inputs = output, fingerprint, inputs
        self.args, self.policies, self.records, self.samples = args, policies, records, samples
        self.output.mkdir(parents=True, exist_ok=True)
        self.saved = None
        final = output/'metrics.json'
        if not args.force and final.exists():
            saved = read_json(final)
            if saved.get('fingerprint') == fingerprint and saved.get('status') == 'complete':
                self.saved = saved
                export_tables(output, saved)
        self.confusions = {k: np.zeros((count, count+1), np.int64)
                           for k in policies if k.startswith('multiclass_')}
        if args.base_dir:
            self.confusions['baseline_multiclass'] = np.zeros((count, count+1), np.int64)
        self.binary = dict(independent_binary=np.zeros((count, 3), np.int64),
                           prompt_binary_ge=np.zeros((len(records), 3), np.int64),
                           prompt_binary_gt=np.zeros((len(records), 3), np.int64))
        self.next_image, self.elapsed = 0, 0.
        self.disagreements = dict(multiclass_saved=0, multiclass_common=0)
        self.refiner_changes = {}
        progress_path = output/'progress.json'
        if not args.force and not self.saved and progress_path.exists():
            progress = read_json(progress_path)
            if progress.get('fingerprint') == fingerprint:
                self.next_image = progress['processed_images']
                if not 0 <= self.next_image <= len(samples):
                    raise ValueError('Invalid progress image count')
                self.confusions = {k: np.asarray(v, dtype=np.int64) for k, v in progress['confusions'].items()}
                self.binary = {k: np.asarray(v, dtype=np.int64) for k, v in progress['binary'].items()}
                self.disagreements = progress['baseline_disagreements']
                self.refiner_changes = progress.get('refiner_changes', {})
                self.elapsed = progress['elapsed_seconds']

    def observe(self, scores, prompt_scores, target, valid, baseline, raw_scores=None):
        count = len(self.confusions['multiclass_common'])
        if scores.shape != (count, *target.shape):
            raise ValueError('Head/GT shape mismatch')
        if baseline is not None:
            add_confusion(self.confusions['baseline_multiclass'], baseline, target, valid)
        for key in ('multiclass_saved', 'multiclass_common'):
            policy = self.policies[key]
            prediction = classify_scores(scores, list(range(count)), policy['threshold'], policy['background'])
            add_confusion(self.confusions[key], prediction, target, valid)
            if raw_scores is not None:
                raw_prediction = classify_scores(raw_scores, list(range(count)), policy['threshold'], policy['background'])
                changes = self.refiner_changes.setdefault(key, dict(changed_pixels=0, corrected_pixels=0, damaged_pixels=0))
                for name, value in prediction_changes(raw_prediction, prediction, target, valid).items():
                    changes[name] += value
            if baseline is not None:
                self.disagreements[key] += int(np.count_nonzero((prediction != baseline) & valid))
        threshold = self.policies['independent_binary']['threshold']
        if raw_scores is not None and 'refined_binary' not in self.binary:
            self.binary['refined_binary'] = np.zeros_like(self.binary['independent_binary'])
        for channel in range(count):
            add_binary(self.binary['independent_binary'], channel,
                       (scores if raw_scores is None else raw_scores)[channel] >= threshold, target == channel, valid)
            if raw_scores is not None:
                add_binary(self.binary['refined_binary'], channel, scores[channel] >= threshold, target == channel, valid)
        for index, record in enumerate(self.records):
            score = prompt_scores[record['prompt_index']]
            if not np.isfinite(score).all() or np.any((score < 0) | (score > 1)):
                raise ValueError(f'Invalid probabilities for prompt {record["prompt"]}')
            truth = target == record['concept_index']
            add_binary(self.binary['prompt_binary_ge'], index, score >= threshold, truth, valid)
            add_binary(self.binary['prompt_binary_gt'], index, score > threshold, truth, valid)

    def save_progress(self):
        atomic_json_save(dict(fingerprint=self.fingerprint, processed_images=self.next_image,
                              elapsed_seconds=self.elapsed,
                              confusions={k: v.tolist() for k, v in self.confusions.items()},
                              binary={k: v.tolist() for k, v in self.binary.items()},
                              baseline_disagreements=self.disagreements,
                              **({'refiner_changes': self.refiner_changes} if self.refiner_changes else {})), self.output/'progress.json')

    def finish(self, common):
        if self.saved is not None:
            return dict(status='skipped_existing', output=str(self.output), result=self.saved)
        metrics = {k: summarize_confusion(v, common['class_names']) for k, v in self.confusions.items()}
        metrics.update({k: summarize_counts(v, [r['prompt'] for r in self.records]
                                           if k.startswith('prompt_') else common['class_names'])
                        for k, v in self.binary.items()})
        result = dict(common, status='complete', version=VERSION, fingerprint=self.fingerprint,
                      finished_at=datetime.now().isoformat(timespec='seconds'), inputs=self.inputs,
                      elapsed_seconds=self.elapsed, policies=self.policies, metrics=metrics)
        if self.refiner_changes:
            result['refiner_changes'] = self.refiner_changes
        if self.args.base_dir:
            baseline_summary = metrics['baseline_multiclass']['summary']
            result['baseline_comparison'] = dict(
                directory=str(self.args.base_dir.resolve()), scope='same evaluated images and valid GT pixels',
                disagreement_valid_pixels=self.disagreements,
                delta_percentage_points={key: metric_delta(metrics[key]['summary'], baseline_summary)
                                         for key in ('multiclass_saved', 'multiclass_common')})
        export_tables(self.output, result)
        atomic_json_save(result, self.output/'metrics.json')
        return dict(status='evaluated', output=str(self.output), result=result)


def metric_delta(current, reference):
    return {key: None if value is None or reference.get(key) is None else value-reference[key]
            for key, value in current.items()}


def evaluate_run(directory, args, candidates=None):
    """Evaluate one best assembly, or many replacements sharing each image's SAM heads.

    Candidates have id, overrides and output fields. Only affected concepts are
    recomputed; all unchanged concept/prompt maps come from the original best.
    """
    device = torch.device(args.device)
    run, concepts, selected, records, prompts, aggregators = load_run(directory, device, **evaluation_refiner_options(args))
    config, cx = run['config'], run['config']['cx_b']
    if selected != list(range(len(concepts))):
        raise ValueError('Final multiclass evaluation requires all concept checkpoints; this run selected a subset')
    if args.common_background not in selected:
        raise ValueError('--common-background must be a defined zero-based concept index')
    saved_background = cx.get('prediction_background', NODATA)
    if saved_background not in selected+[NODATA]:
        raise ValueError('Saved prediction_background is not a concept index or 65535')
    names = [c.get('concept_name') or c['concept_id'] for c in concepts]
    is_json = str(config.get('concept_path', config['model']['classname_path'])).lower().endswith('.json')
    if args.base_dir and (is_json or [c['concept_id'] for c in concepts] != LOVEDA_CLASSES):
        raise ValueError('--base-dir comparison requires the seven original LoveDA concepts in order')
    all_samples, dataset = collect_samples(config, args)
    samples = all_samples if args.max_images == -1 else all_samples[:args.max_images]
    ignore = dataset.get('ignore_index', 255) if is_json else 0
    threshold = cx.get('metric_threshold', .5)
    policies = dict(
        multiclass_saved=dict(threshold=config['model'].get('prob_thd', .1), background=saved_background),
        multiclass_common=dict(threshold=args.common_threshold, background=args.common_background),
        independent_binary=dict(threshold=threshold, comparison='>='),
        prompt_binary_ge=dict(threshold=threshold, comparison='>='),
        prompt_binary_gt=dict(threshold=threshold, comparison='>'))
    fingerprint, inputs = evaluation_signature(directory, run, samples, args, policies)
    single = candidates is None
    candidates = candidates if candidates is not None else [dict(id='original_best', overrides={}, output=output_directory(directory, args))]
    states, modules = {}, {}
    for candidate in candidates:
        overrides = candidate['overrides']
        signature, source = candidate_signature(inputs, directory, overrides) if overrides else (fingerprint, inputs)
        joint_checkpoint = candidate.get('joint_checkpoint')
        if joint_checkpoint:
            source = dict(inputs, checkpoint_sha256={joint_checkpoint: sha256(directory/joint_checkpoint)},
                          joint_checkpoint=joint_checkpoint)
            signature = hashlib.sha256(json.dumps(source, sort_keys=True).encode()).hexdigest()
        state = CandidateEvaluation(candidate['output'], signature, source, len(concepts), records, samples, args, policies)
        states[candidate['id']] = state
        if (overrides or joint_checkpoint) and state.saved is None and state.next_image < len(samples):
            loaded = (load_run(directory, device, joint_checkpoint=joint_checkpoint, **evaluation_refiner_options(args)) if joint_checkpoint
                      else load_run(directory, device, overrides, **evaluation_refiner_options(args)))
            changed = set()
            if joint_checkpoint:
                changed.update(selected)
            for slot in overrides:
                if slot.startswith('concept/'):
                    changed.add(int(Path(slot).parent.name.split('_')[1]))
                else:
                    prompt_index = int(Path(slot).parent.name.split('_')[1])
                    changed.add(next(r['concept_index'] for r in records if r['prompt_index'] == prompt_index))
            modules[candidate['id']] = (loaded[4], loaded[5], sorted(changed))
    pending = [s for s in states.values() if s.saved is None and s.next_image < len(samples)]
    next_image = min((s.next_image for s in pending), default=len(samples))
    started = time.monotonic()
    print(f'[{directory.parent.name}/{directory.name}] {len(samples)}/{len(all_samples)} images; '
          f'candidates={len(candidates)} pending={len(pending)} resume={next_image}; mode={cx["data_mode"]}', flush=True)
    cache, provider = (None, None) if not pending else build_head_source(run, records, args, device)
    with torch.inference_mode():
        for index in range(next_image, len(samples)):
            image_path, gt_path, relative = samples[index]
            with Image.open(gt_path) as image:
                gt = np.asarray(image).copy()
            if gt.ndim != 2:
                raise ValueError(f'GT must be single-channel: {gt_path}')
            with Image.open(image_path) as image:
                if (image.height, image.width) != gt.shape:
                    raise ValueError(f'Image/GT shape mismatch: {image_path}')
            target = map_gt(gt, concepts, is_json, ignore)
            valid = gt != ignore
            if np.any(valid & (target == NODATA)):
                raise ValueError(f'Nonignored GT labels are not covered by the concepts: {gt_path}')
            online = provider.image_outputs(image_path, [r['prompt_index'] for r in records]) if provider else None
            # Load each raw head once per image, regardless of candidate count.
            heads = ({r['prompt_index']: online[r['prompt_index']]['heads'] for r in records}
                     if online is not None else {r['prompt_index']: cache.load(r, index, image_path) for r in records})
            get_heads = lambda r: heads[r['prompt_index']]
            scores, prompt_scores = predict_image(records, selected, prompts, aggregators, get_heads, online)
            baseline = None
            if args.base_dir:
                with Image.open(args.base_dir/'result'/image_path.name) as image:
                    baseline = np.asarray(image).astype(np.int64)-1
                if baseline.shape != gt.shape or np.any((baseline < 0) | (baseline >= len(concepts))):
                    raise ValueError(f'Expected baseline PNG with original LoveDA IDs 1..7: {image_path.name}')
            processed = index+1
            for candidate in candidates:
                state = states[candidate['id']]
                if state.saved is not None or state.next_image > index:
                    continue
                candidate_started = time.monotonic()
                candidate_scores, candidate_prompts = scores, prompt_scores
                candidate_raw = getattr(aggregators, 'raw_scores', None)
                if candidate['overrides'] or candidate.get('joint_checkpoint'):
                    candidate_modules, candidate_aggregators, changed = modules[candidate['id']]
                    replacement, replacement_prompts = predict_image(
                        records, changed, candidate_modules, candidate_aggregators, get_heads, online)
                    candidate_scores = scores.copy()
                    candidate_scores[changed] = replacement
                    candidate_prompts = dict(prompt_scores)
                    candidate_prompts.update(replacement_prompts)
                    candidate_raw = getattr(candidate_aggregators, 'raw_scores', None)
                state.observe(candidate_scores, candidate_prompts, target, valid, baseline, raw_scores=candidate_raw)
                state.next_image = processed
                state.elapsed += time.monotonic()-candidate_started
                if processed % args.save_every == 0 or processed == len(samples):
                    state.save_progress()
            if processed % args.progress_every == 0 or processed == len(samples):
                elapsed = time.monotonic()-started
                eta = elapsed/(processed-next_image)*(len(samples)-processed)
                print(f'  {processed}/{len(samples)} elapsed={elapsed:.0f}s ETA={eta:.0f}s {relative}', flush=True)
    common = dict(source_directory=str(directory), class_names=names,
                  methods={k: cx[k] for k in ('prompt_method','concept_method','loss_method','data_mode')},
                  seed=run.get('seed'), split=args.split, images=len(samples), dataset_images=len(all_samples),
                  evaluation_scope='full' if len(samples)==len(all_samples) else 'partial',
                  scale='percent',
                  gt_ignore_index=ignore, selected_concepts=selected,
                  notes=['multiclass_saved follows eval_cx_b.py, including 65535=unassigned when prediction_background is absent.',
                         'multiclass_common changes only final decision; SAM filtering and weights are unchanged.',
                         'Undefined per-class metrics are null and excluded from their corresponding macro means.',
                         'independent_binary is not final multiclass; no binary aAcc is reported.',
                         'Input cache/data signatures use path, size and mtime; checkpoint/source files use SHA256.',
                         'elapsed_seconds counts candidate computation; shared SAM/head loading is excluded.'])
    results = {key: state.finish(common) for key, state in states.items()}
    return results[candidates[0]['id']] if single else results


def checkpoint_candidates(directory, args):
    """Enumerate every pth in each active module slot; unrelated root pth are not models."""
    run = read_json(directory/'run_config.json')
    folder = 'dynamic_prompt' if run['config']['cx_b']['prompt_method'] == 'dynamic_head' else 'prompt_checkpoints'
    best_paths = sorted((directory/folder).glob('*/best.pth')) + sorted((directory/'concept').glob('*/best.pth'))
    candidates = [dict(id='original_best', overrides={}, output=output_directory(directory, args))]
    inventory = []
    for best in best_paths:
        slot = best.relative_to(directory).as_posix()
        for path in sorted(best.parent.glob('*.pth')):
            relative = path.relative_to(directory).as_posix()
            payload = torch.load(path, map_location='cpu')
            meta = payload.get('metadata', {})
            item = dict(checkpoint=relative, slot=slot, kind=meta.get('kind'),
                        epoch=meta.get('epoch'), checkpoint_type=meta.get('checkpoint_type'),
                        candidate_id='original_best' if path == best else relative)
            inventory.append(item)
            if path != best:
                candidates.append(dict(id=relative, overrides={slot: relative},
                    output=output_directory(directory, args)/'checkpoint_scan'/'candidates'/Path(relative).with_suffix('')))
    return candidates, inventory


def report_row(identifier, evaluated, original, **extra):
    result = evaluated['result']
    metrics = result['metrics']['multiclass_common']['summary']
    reference = original['result']['metrics']['multiclass_common']['summary']
    row = dict(candidate_id=identifier, output=evaluated['output'], images=result['images'],
               evaluation_scope=result['evaluation_scope'], **extra, **metrics)
    row.update({'delta_vs_original_best_'+k: v for k, v in metric_delta(metrics, reference).items()})
    if 'baseline_comparison' in result:
        baseline = result['metrics']['baseline_multiclass']['summary']
        row.update({'loveda_'+k: v for k, v in baseline.items()})
        row.update({'delta_vs_loveda_'+k: v for k, v in metric_delta(metrics, baseline).items()})
    return row


def scan_checkpoints(directory, args):
    if read_json(directory/'run_config.json')['config']['cx_b'].get('joint_mode', 'independent') != 'independent':
        return scan_joint_checkpoints(directory, args)
    candidates, inventory = checkpoint_candidates(directory, args)
    print(f'[checkpoint scan] {len(inventory)} pth; {len(candidates)-1} single-module replacements', flush=True)
    evaluated = evaluate_run(directory, args, candidates)
    original = evaluated['original_best']
    def score(identifier):
        value = evaluated[identifier]['result']['metrics']['multiclass_common']['summary']['mIoU']
        if value is None or not np.isfinite(value):
            raise ValueError(f'Cannot select a checkpoint with undefined final mIoU: {identifier}')
        return value
    score('original_best')
    # Original best wins ties; each slot is compared in the same original context.
    winners = {}
    for item in inventory:
        slot, identifier = item['slot'], item['candidate_id']
        if slot not in winners:
            winners[slot] = 'original_best'
        if score(identifier) > score(winners[slot]):
            winners[slot] = identifier
    overrides = {slot: identifier for slot, identifier in winners.items() if identifier != 'original_best'}
    scan_output = output_directory(directory, args)/'checkpoint_scan'
    assembly = dict(id='winners_assembly', overrides=overrides, output=scan_output/'winners_assembly')
    if overrides:
        evaluated.update(evaluate_run(directory, args, [assembly]))
    else:
        # No improvement in any slot: the proposed assembly is exactly original_best.
        assembly['output'].mkdir(parents=True, exist_ok=True)
        export_tables(assembly['output'], original['result'])
        atomic_json_save(original['result'], assembly['output']/'metrics.json')
        evaluated['winners_assembly'] = dict(original, output=str(assembly['output']))
    candidates.append(assembly)
    best_single = max((c['id'] for c in candidates if c['id'] != 'winners_assembly'), key=score)
    selected_id = max((c['id'] for c in candidates), key=score)
    selected_spec = next(c for c in candidates if c['id'] == selected_id)
    baseline_manifest = {slot: slot for slot in winners}
    manifest = dict(baseline_manifest, **selected_spec['overrides'])
    assembly_manifest = dict(baseline_manifest, **overrides)
    rows = []
    for item in inventory:
        identifier = item['candidate_id']
        details = {k: v for k, v in item.items() if k != 'candidate_id'}
        if identifier == 'original_best':
            # Every pth, including each best.pth, gets its own exported metrics.
            output = scan_output/'candidates'/Path(item['checkpoint']).with_suffix('')
            output.mkdir(parents=True, exist_ok=True)
            export_tables(output, original['result'])
            atomic_json_save(original['result'], output/'metrics.json')
            value = dict(original, output=str(output))
        else:
            value = evaluated[identifier]
        rows.append(report_row(identifier, value, original, **details))
    assembly_rows = [report_row(identifier, evaluated[identifier], original)
                     for identifier in dict.fromkeys(['original_best', best_single, 'winners_assembly', selected_id])]
    report = dict(status='complete', selection_metric='multiclass_common.mIoU',
                  search_strategy='one module replaced at a time; all other modules use original best.pth',
                  scope='Best among evaluated candidates, not exhaustive Cartesian combinations.',
                  images=original['result']['images'], evaluation_scope=original['result']['evaluation_scope'],
                  original_best=report_row('original_best', original, original),
                  best_single_checkpoint=report_row(best_single, evaluated[best_single], original),
                  winners_assembly=report_row('winners_assembly', evaluated['winners_assembly'], original),
                  selected=report_row(selected_id, evaluated[selected_id], original),
                  module_winners={slot: (slot if value == 'original_best' else value) for slot, value in winners.items()},
                  selected_checkpoint_manifest=manifest, winners_assembly_manifest=assembly_manifest,
                  selected_checkpoint_sha256={path: sha256(directory/path) for path in manifest.values()},
                  checkpoints=rows, assemblies=assembly_rows,
                  notes=['All rankings and deltas use the same evaluated images and final multiclass_common policy.',
                         'Module winners can interact; their assembly is evaluated separately and may be worse.',
                         'Only saved checkpoints are searched; unsaved epochs cannot be reconstructed.',
                         'No training checkpoints are overwritten.'])
    scan_output.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for name, table in [('checkpoints', rows), ('assemblies', assembly_rows)]:
        headers = list(dict.fromkeys(k for row in table for k in row))
        with (scan_output/(name+'.csv')).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            writer.writerows(table)
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in table:
            sheet.append([row.get(k) for k in headers])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    book.save(scan_output/'checkpoint_comparison.tmp.xlsx')
    book.close()
    os.replace(scan_output/'checkpoint_comparison.tmp.xlsx', scan_output/'checkpoint_comparison.xlsx')
    atomic_json_save(report, scan_output/'checkpoint_comparison.json')
    atomic_json_save(dict(source_directory=str(directory), selected_candidate=selected_id,
                          selection_metric=report['selection_metric'], images=report['images'],
                          evaluation_scope=report['evaluation_scope'], checkpoints=manifest,
                          sha256=report['selected_checkpoint_sha256']), scan_output/'selected_checkpoints.json')
    def display(value):
        return '—' if value is None else f'{value:.6f}'
    lines = ['# Checkpoint 多分类评估', '',
             f"评估范围：{report['evaluation_scope']}，图片数：{report['images']}。",
             '按最终 multiclass_common mIoU 选优；差值单位为百分点。', '',
             '| 候选 | mIoU (%) | 相对原 best | LoveDA mIoU (%) | 相对 LoveDA |',
             '| --- | ---: | ---: | ---: | ---: |']
    for row in assembly_rows:
        lines.append(f"| {row['candidate_id']} | {display(row['mIoU'])} | "
                     f"{display(row['delta_vs_original_best_mIoU'])} | {display(row.get('loveda_mIoU'))} | "
                     f"{display(row.get('delta_vs_loveda_mIoU'))} |")
    lines += ['', f'评估候选中最优：`{selected_id}`。',
              '完整权重清单见 `selected_checkpoints.json`；每个 pth 的指标和差值见 `checkpoints.csv`。',
              '各模块优胜权重组装后已单独评估；这里未穷举所有模块的笛卡尔积，不代表全局最优。',
              '仅评估已保存的 pth，无法补回未保存的 epoch；原训练权重保持不变。', '']
    (scan_output/'summary.md').write_text('\n'.join(lines), encoding='utf-8')
    print(f"[checkpoint scan] original={score('original_best'):.6f}; best single={score(best_single):.6f}; "
          f"assembled={score('winners_assembly'):.6f}; selected={selected_id}", flush=True)
    output = dict(original)
    output['status'] = 'skipped_existing' if all(v['status']=='skipped_existing' for v in evaluated.values()) else 'evaluated'
    output.update(checkpoint_scan_output=str(scan_output), checkpoint_count=len(inventory),
                  selected_candidate=selected_id, selected_checkpoints=str(scan_output/'selected_checkpoints.json'))
    for label, value in [('best_single', report['best_single_checkpoint']),
                         ('winners_assembly', report['winners_assembly']), ('selected', report['selected'])]:
        output.update({label+'_'+key: val for key, val in value.items()
                       if key.startswith(('m', 'aAcc', 'delta_vs_', 'loveda_'))})
    return output


def scan_joint_checkpoints(directory, args):
    """Joint epochs are whole pipelines; never combine their independently saved modules."""
    scan_output = output_directory(directory, args)/'checkpoint_scan'
    candidates = [dict(id='original_best', overrides={}, output=output_directory(directory, args))]
    inventory = []
    for path in sorted((directory/'joint').glob('*.pth')):
        payload = torch.load(path, map_location='cpu')
        if payload.get('format') not in JOINT_FORMATS:
            raise ValueError(f'Invalid joint checkpoint: {path}')
        relative = path.relative_to(directory).as_posix()
        inventory.append(dict(checkpoint=relative, epoch=payload['epoch']))
        candidates.append(dict(id=relative, overrides={}, joint_checkpoint=relative,
                               output=scan_output/'candidates'/path.stem))
    if not inventory:
        raise ValueError('Joint run has no complete pipeline checkpoints')
    evaluated = evaluate_run(directory, args, candidates)
    original = evaluated['original_best']
    rows = [report_row(item['checkpoint'], evaluated[item['checkpoint']], original, **item) for item in inventory]
    original_row = report_row('original_best', original, original)
    if any(row['mIoU'] is None for row in [original_row]+rows):
        raise ValueError('Cannot select joint checkpoints without valid multiclass mIoU')
    selected = max([original_row]+rows, key=lambda row: row['mIoU'])
    selected_path = selected.get('checkpoint', 'joint/best.pth')
    scan_output.mkdir(parents=True, exist_ok=True)
    report = dict(status='complete', selection_metric='multiclass_common.mIoU',
                  search_strategy='complete joint epoch checkpoints', original_best=original_row,
                  selected=selected, checkpoints=rows,
                  selected_joint_checkpoint=selected_path, sha256=sha256(directory/selected_path))
    atomic_json_save(report, scan_output/'checkpoint_comparison.json')
    atomic_json_save(dict(source_directory=str(directory), joint_checkpoint=selected_path,
                          sha256=report['sha256']), scan_output/'selected_checkpoints.json')
    headers = list(dict.fromkeys(k for row in rows for k in row))
    with (scan_output/'checkpoints.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    book = Workbook()
    sheet = book.active
    sheet.title = 'checkpoints'
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(k) for k in headers])
    book.save(scan_output/'checkpoint_comparison.xlsx')
    book.close()
    (scan_output/'summary.md').write_text(
        '# 联调完整模型 checkpoint 比较\n\n'
        f"范围：{selected['evaluation_scope']}，图片数：{selected['images']}。\n\n"
        f"原 best mIoU：{original_row['mIoU']:.6f}%\n\n"
        f"最优：`{selected_path}`，mIoU：{selected['mIoU']:.6f}%\n\n"
        f"相对原 best：{selected['delta_vs_original_best_mIoU']:.6f} 个百分点。\n\n"
        f"相对 LoveDA：{selected.get('delta_vs_loveda_mIoU')} 个百分点。\n\n"
        '按完整联调 checkpoint 评估，不混合不同 epoch 的模块。详见 checkpoints.csv。\n', encoding='utf-8')
    output = dict(original, checkpoint_scan_output=str(scan_output), checkpoint_count=len(rows),
                  selected_candidate=selected_path, selected_checkpoints=str(scan_output/'selected_checkpoints.json'))
    output['status'] = 'skipped_existing' if all(v['status']=='skipped_existing' for v in evaluated.values()) else 'evaluated'
    output.update({'selected_'+key: value for key, value in selected.items()
                   if key.startswith(('m', 'aAcc', 'delta_vs_', 'loveda_'))})
    return output


def main(argv=None):
    args = parse_args(argv)
    torch.set_num_threads(args.cpu_threads)
    entries = discover_runs(args)
    if args.list_only:
        print(json.dumps(entries, ensure_ascii=False, indent=2))
        return 1 if any(r['status']=='failed' for r in entries) else 0
    for row in entries:
        if row['status'] != 'ready':
            print(f'[{row["status"]}] {row["directory"]}: {row.get("error", "training not complete")}', flush=True)
            continue
        directory = Path(row['directory'])
        try:
            evaluated = evaluate_run(directory, args) if args.use_best else scan_checkpoints(directory, args)
            result = evaluated.pop('result')
            row.update(evaluated, images=result['images'], evaluation_scope=result['evaluation_scope'])
            for policy in ('multiclass_saved', 'multiclass_common', 'independent_binary'):
                row.update({policy+'_'+k: v for k, v in result['metrics'][policy]['summary'].items()})
            if 'baseline_comparison' in result:
                row.update({'loveda_'+k: v for k, v in result['metrics']['baseline_multiclass']['summary'].items()})
                for policy, deltas in result['baseline_comparison']['delta_percentage_points'].items():
                    row.update({policy+'_delta_vs_loveda_'+k: v for k, v in deltas.items()})
            print(f'[{row["status"]}] {directory.parent.name}: common mIoU={row["multiclass_common_mIoU"]}', flush=True)
        except Exception as error:
            row.update(status='failed', error=f'{type(error).__name__}: {error}')
            output = output_directory(directory, args)
            output.mkdir(parents=True, exist_ok=True)
            atomic_json_save(dict(status='failed', error=row['error'], traceback=traceback.format_exc(),
                                  failed_at=datetime.now().isoformat(timespec='seconds')), output/'last_error.json')
            print(f'[failed] {directory}: {row["error"]}', flush=True)
        finally:
            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()
    root = args.out_root or args.root
    root.mkdir(parents=True, exist_ok=True)
    suffix = args.split if args.max_images == -1 else f'{args.split}_first_{args.max_images}'
    if not args.use_best:
        suffix += '_all_checkpoints'
    if ablation_folder(args):
        suffix += '__'+ablation_folder(args)
    atomic_json_save(dict(created_at=datetime.now().isoformat(timespec='seconds'), runs=entries),
                     root/f'batch_multiclass_{suffix}.json')
    headers = list(dict.fromkeys(key for row in entries for key in row))
    with (root/f'batch_multiclass_{suffix}.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(entries)
    print(f'Batch summary: {root/f"batch_multiclass_{suffix}.csv"}', flush=True)
    return 1 if any(r['status']=='failed' for r in entries) else 0


if __name__ == '__main__':
    raise SystemExit(main())
