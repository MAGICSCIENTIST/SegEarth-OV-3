import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path

import torch
from openpyxl import load_workbook

from cx_b.checkpoints import atomic_torch_save, completed_matches
from cx_b.configuration import (
    COMBINATION_CONFIGS, ConfigurationError, apply_dotted_overrides,
    config_fingerprint, resolve_combination_config, validate_config,
)
from cx_b.heads import export_sam_region_cache, validate_cache_contract
from cx_b.losses import IndependentBinaryLoss, build_binary_loss
from cx_b.modules import (
    A3PromptAdapter, DynamicPromptHead, SegOV3PromptAdapter,
    build_prompt_module, build_concept_aggregator,
    load_variant, region_support, save_variant,
)
from metrics_reporting import (
    MetricsReporter, ProgressReporter, TrainingOutputReporter,
)
from cx_b.trainer import (
    CXBTrainer, DEFAULT_A3_PARAMETERS, optimize_binary_module,
    should_reuse_a3_prompt,
)
from train_batch_prompt_concept_variants import parse_args


class CXBTests(unittest.TestCase):
    def test_all_configs_load_with_exact_identity(self):
        self.assertEqual(len(COMBINATION_CONFIGS), 105)
        for identity in COMBINATION_CONFIGS:
            cfg, path = resolve_combination_config(*identity)
            self.assertEqual((cfg['cx_b']['prompt_method'],
                              cfg['cx_b']['concept_method'],
                              cfg['cx_b']['loss_method']), identity)
            self.assertEqual(cfg['cx_b']['prompt_init'], 'scratch')
            self.assertFalse(cfg['cx_b']['export_probabilities'])
            self.assertEqual(
                cfg['cx_b']['prompt_train']['checkpoint_interval_epochs'], 10)
            self.assertEqual(
                cfg['cx_b']['concept_train']['checkpoint_interval_epochs'], 10)
            self.assertTrue(Path(path).is_file())

    def test_cli_supports_documented_typo_and_alias(self):
        a = parse_args(['--prompt_methd=A3-2', '--concept_method=max',
                        '--loss_method=A3-2', '--dry-run'])
        b = parse_args(['--prompt_method=A3-2', '--concept_method=max',
                        '--loss_method=A3-2', '--dry-run'])
        self.assertEqual(a.prompt_method, b.prompt_method)
        c = parse_args(['--prompt_method=A3-2', '--concept_method=max',
                        '--loss_method=A3-2', '--prompt_init=a3_warmstart'])
        self.assertEqual(c.prompt_init, 'a3_warmstart')
        d = parse_args(['--prompt_method=A3-2', '--concept_method=max',
                        '--loss_method=A3-2',
                        '--export-probabilities-only'])
        self.assertTrue(d.export_probabilities_only)

    def test_prompt_init_validation_and_reuse_are_explicit(self):
        config = dict(cx_b=dict(
            prompt_method='A3-2', concept_method='max', loss_method='A3-2',
            prompt_init='invalid', concept_independent=True, data_mode='cache',
            sam_region_cfg={}))
        with self.assertRaises(ConfigurationError):
            validate_config(config)
        self.assertTrue(should_reuse_a3_prompt(
            'a3_reuse', 'A3-2', 'A3-2'))
        self.assertFalse(should_reuse_a3_prompt(
            'scratch', 'A3-2', 'A3-2'))
        self.assertFalse(should_reuse_a3_prompt(
            'a3_warmstart', 'A3-2', 'A3-2'))

    def test_scratch_initialization_does_not_read_a3(self):
        expected = dict(
            prompt_index=0, prompt='building', concept_index=0,
            concept_id='1', concept_name='building', prompt_position=0,
            valid_labels=[1], prompt_score_version='identifiable_three_v1')
        with tempfile.TemporaryDirectory() as directory:
            trainer = CXBTrainer(
                {'cx_b': {'prompt_init': 'scratch'}}, __file__, directory)
            records = trainer._load_base_records([expected])
            self.assertEqual(records[0]['parameters'], DEFAULT_A3_PARAMETERS)
            with self.assertRaises(ValueError):
                CXBTrainer(
                    {'cx_b': {'prompt_init': 'scratch'}}, __file__, directory,
                    a3_checkpoint_root='old_a3')

    def test_warmstart_loads_a_complete_a3_source(self):
        expected = dict(
            prompt_index=0, prompt='building', concept_index=0,
            concept_id='1', concept_name='building', prompt_position=0,
            valid_labels=[1], prompt_score_version='identifiable_three_v1')
        parameters = dict(
            semantic_weight=.72, presence_gain=1.4,
            presence_intercept=-.2)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / 'a3_source'
            checkpoint_dir = root / 'prompt_best_params'
            checkpoint_dir.mkdir(parents=True)
            atomic_torch_save(
                {'record': dict(expected, parameters=parameters)},
                checkpoint_dir / 'prompt_000.pth')
            trainer = CXBTrainer(
                {'cx_b': {'prompt_init': 'a3_warmstart'}}, __file__,
                Path(directory) / 'output', a3_checkpoint_root=root)
            records = trainer._load_base_records([expected])
            self.assertEqual(records[0]['parameters'], parameters)

    def test_protected_override_is_rejected(self):
        with self.assertRaises(ConfigurationError):
            apply_dotted_overrides(
                {'cx_b': {}}, {'cx_b.concept_method': 'max'})

    def test_a3_formula_and_dynamic_zero_initialization_match(self):
        heads = dict(
            semantic_score_map=torch.rand(8, 9),
            instance_score_map=torch.rand(8, 9),
            presence_score=torch.tensor(0.8))
        base = A3PromptAdapter(.63, 1.2, -.1)
        dynamic = DynamicPromptHead(.63, 1.2, -.1)
        self.assertTrue(torch.allclose(base(heads), dynamic(heads), atol=1e-6))

    def test_scratch_dynamic_head_trains_base_and_dynamic_branch(self):
        module = DynamicPromptHead(.5, 1., 0., base_trainable=True)
        heads = dict(
            semantic_score_map=torch.rand(8, 9),
            instance_score_map=torch.rand(8, 9),
            presence_score=torch.tensor(.8))
        module(heads).mean().backward()
        self.assertTrue(all(
            parameter.requires_grad for parameter in module.base.parameters()))
        self.assertTrue(any(
            parameter.grad is not None
            for parameter in module.base.parameters()))
        self.assertIsNotNone(module.output_projection.weight.grad)

    def test_every_concept_method_is_exact_identity_for_q1(self):
        values = torch.tensor([[[0., 1.], [.3, .7]]], dtype=torch.float64)
        for method in ('A3-2', 'max', 'logsumexp', 'dynamic',
                       'reliability_dynamic', 'region_dynamic',
                       'sam_region_dynamic'):
            module = build_concept_aggregator(method, 1)
            output = module(values)
            self.assertEqual(output.dtype, torch.float64)
            self.assertTrue(torch.equal(output, values[0]), method)

    def test_multi_prompt_weights_normalize_and_output_is_convex(self):
        values = torch.rand(3, 8, 9)
        for method in ('A3-2', 'max', 'dynamic', 'reliability_dynamic',
                       'region_dynamic', 'sam_region_dynamic'):
            module = build_concept_aggregator(method, 3)
            output, weights = module(values, return_weights=True)
            self.assertTrue(torch.allclose(weights.sum(0),
                                           torch.ones(8, 9), atol=1e-6))
            self.assertTrue(torch.all(output >= values.min(0).values - 1e-6))
            self.assertTrue(torch.all(output <= values.max(0).values + 1e-6))

    def test_logsumexp_has_positive_temperature_and_normalized_weights(self):
        module = build_concept_aggregator('logsumexp', 3, temperature=.2)
        output, weights = module(torch.rand(3, 4, 5), return_weights=True)
        self.assertGreater(float(module.temperature), 0)
        self.assertTrue(torch.allclose(weights.sum(0), torch.ones(4, 5)))

    def test_binary_losses_handle_empty_positive_and_all_ignore(self):
        prediction = torch.full((3, 4), .8, requires_grad=True)
        target = torch.zeros_like(prediction)
        for method in ('A3-2', 'soft_iou_bce', 'bce_tversky'):
            loss = IndependentBinaryLoss(method)(
                prediction, target, torch.ones_like(target, dtype=torch.bool))
            self.assertIsNotNone(loss)
            self.assertGreater(float(loss), 0)
            self.assertIsNone(IndependentBinaryLoss(method)(
                prediction, target, torch.zeros_like(target, dtype=torch.bool)))

    def test_eval_style_summary_contains_all_requested_metrics(self):
        rows = [
            dict(IoU=50., Acc=80., Fscore=60., Precision=50., Recall=80.,
                 InitialIoU=40., InitialAcc=70., InitialFscore=50.,
                 InitialPrecision=40., InitialRecall=70., InitialTP=7,
                 TP=8, GTPixels=10),
            dict(IoU=25., Acc=50., Fscore=40., Precision=33., Recall=50.,
                 InitialIoU=20., InitialAcc=40., InitialFscore=30.,
                 InitialPrecision=25., InitialRecall=40., InitialTP=4,
                 TP=5, GTPixels=10),
        ]
        summary = MetricsReporter.summarize_rows(
            rows, epoch=3, split='validation', loss=.2)
        self.assertEqual(set((
            'aAcc', 'mIoU', 'mAcc', 'mFscore', 'mPrecision', 'mRecall'))
            - set(summary), set())
        self.assertAlmostEqual(summary['aAcc'], 65.)
        self.assertAlmostEqual(summary['mIoU'], 37.5)
        self.assertAlmostEqual(summary['mFscore'], 50.)
        self.assertAlmostEqual(summary['InitialmIoU'], 30.)
        self.assertAlmostEqual(summary['mIoUGain'], 7.5)
        self.assertAlmostEqual(summary['InitialaAcc'], 55.)
        self.assertAlmostEqual(summary['aAccGain'], 10.)
        self.assertAlmostEqual(summary['mAccGain'], 10.)
        self.assertAlmostEqual(summary['mFscoreGain'], 10.)
        self.assertAlmostEqual(summary['mPrecisionGain'], 9.)
        self.assertAlmostEqual(summary['mRecallGain'], 10.)

    def test_eval_style_workbook_has_summary_concept_and_prompt_sheets(self):
        class_rows = [dict(
            Class='building', IoU=50., Acc=80., Fscore=60., Precision=50.,
            Recall=80., TP=8, FP=8, FN=2, PredPixels=16, GTPixels=10,
            PredGTAreaRatio=1.6, InitialIoU=40., IoUGain=10.,
            InitialAcc=70., AccGain=10., InitialFscore=50., FscoreGain=10.,
            InitialPrecision=40., PrecisionGain=10., InitialRecall=70.,
            RecallGain=10., InitialTP=7, BestEpoch=6,
            MaxBaselineIoU=45., IoUDeltaVsMax=5.)]
        prompt_rows = [dict(
            Prompt='house', ConceptID='building', ConceptName='building',
            IoU=40., Acc=60., Fscore=50., Precision=43., Recall=60.,
            TP=6, FP=8, FN=4, PredPixels=14, GTPixels=10,
            PredGTAreaRatio=1.4, InitialIoU=35., IoUGain=5.,
            InitialAcc=50., AccGain=10., InitialFscore=40., FscoreGain=10.,
            InitialPrecision=33., PrecisionGain=10., InitialRecall=50.,
            RecallGain=10., InitialTP=5, BestEpoch=4.)]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'results.xlsx'
            MetricsReporter.write_workbook(
                path,
                summaries=[MetricsReporter.summarize_rows(class_rows)],
                class_metrics=class_rows,
                prompt_metrics=prompt_rows,
                model='CX-B', dataset='LoveDA')
            workbook = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(
                set(('summary', 'per_class', 'per_prompt'))
                - set(workbook.sheetnames), set())
            summary_headers = set(next(workbook['summary'].values))
            self.assertEqual(set((
                'aAcc', 'mIoU', 'mAcc', 'mFscore', 'mPrecision',
                'mRecall', 'InitialaAcc', 'aAccGain',
                'InitialmIoU', 'mIoUGain', 'InitialmAcc', 'mAccGain',
                'InitialmFscore', 'mFscoreGain', 'InitialmPrecision',
                'mPrecisionGain', 'InitialmRecall', 'mRecallGain'))
                - summary_headers, set())
            class_headers = next(workbook['per_class'].values)
            prompt_headers = next(workbook['per_prompt'].values)
            self.assertIn('Fscore', class_headers)
            self.assertIn('InitialIoU', class_headers)
            self.assertIn('IoUDeltaVsMax', class_headers)
            self.assertIn('PrecisionGain', class_headers)
            self.assertIn('BestEpoch', class_headers)
            self.assertIn('Fscore', prompt_headers)
            self.assertIn('IoUGain', prompt_headers)
            self.assertIn('RecallGain', prompt_headers)
            self.assertIn('BestEpoch', prompt_headers)
            workbook.close()

    def test_shared_progress_reports_eta_and_expected_finish(self):
        progress = ProgressReporter(total=4, started_at=100.)
        state = progress.snapshot(1, now=110.)
        self.assertEqual(state['elapsed_seconds'], 10.)
        self.assertEqual(state['eta_seconds'], 30.)
        self.assertEqual(state['expected_finish_timestamp'], 140.)
        text = progress.format(1, now=110.)
        self.assertIn('elapsed=00:10', text)
        self.assertIn('eta=00:30', text)
        self.assertIn('expected_finish=', text)

    def test_shared_training_output_mirrors_both_logs(self):
        with tempfile.TemporaryDirectory() as directory:
            with TrainingOutputReporter(directory, 'unit training'):
                print('shared-output-marker')
            for name in ('log.txt', 'run.log'):
                content = (Path(directory) / name).read_text(encoding='utf-8')
                self.assertIn('unit training started', content)
                self.assertIn('shared-output-marker', content)
                self.assertIn('expected_finish=', content)

    def test_independent_loss_does_not_touch_other_concept(self):
        current = torch.rand(4, 5, requires_grad=True)
        other = torch.rand(4, 5, requires_grad=True)
        loss = IndependentBinaryLoss('A3-2')(
            current, torch.ones_like(current),
            torch.ones_like(current, dtype=torch.bool))
        loss.backward()
        self.assertIsNotNone(current.grad)
        self.assertIsNone(other.grad)

    def test_region_support_handles_overlap_and_empty_regions(self):
        probabilities = torch.rand(2, 3, 3)
        masks = torch.tensor([[[1, 1, 0], [0, 0, 0], [0, 0, 0]],
                              [[0, 1, 1], [0, 0, 0], [0, 0, 0]]]).bool()
        support = region_support(probabilities, [
            (masks, torch.tensor([.4, .8])),
            (torch.empty(0, 3, 3), torch.empty(0))])
        self.assertEqual(tuple(support.shape), (2, 3, 3))
        self.assertAlmostEqual(float(support[0, 0, 1]), .6, places=6)
        self.assertEqual(float(support[1].sum()), 0.)

    def test_online_contract_rejects_disk_cache_flags(self):
        with self.assertRaises(ValueError):
            validate_cache_contract('online_sam', {'read_output_cache': True})
        self.assertIsNone(export_sam_region_cache({}, {}, {}))

    def test_dynamic_checkpoint_round_trip_prediction(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'best.pth'
            metadata = dict(kind='prompt', method='dynamic_head',
                            base_parameters=dict(semantic_weight=.4,
                                                 presence_gain=1.3,
                                                  presence_intercept=.2),
                            hidden_channels=16, base_trainable=True)
            module = DynamicPromptHead(.4, 1.3, .2, base_trainable=True)
            save_variant(path, module, metadata)
            loaded, loaded_metadata = load_variant(path)
            heads = dict(semantic_score_map=torch.rand(4, 5),
                         instance_score_map=torch.rand(4, 5),
                         presence_score=torch.tensor(.7))
            self.assertTrue(torch.allclose(module(heads), loaded(heads)))
            self.assertEqual(loaded_metadata['method'], 'dynamic_head')
            self.assertTrue(all(
                parameter.requires_grad for parameter in loaded.base.parameters()))

    def test_atomic_completion_marker_requires_same_fingerprint(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'completed.json'
            path.write_text(json.dumps({'fingerprint': 'abc'}), encoding='utf-8')
            self.assertTrue(completed_matches(path, 'abc'))
            self.assertFalse(completed_matches(path, 'changed'))

    def test_run_fingerprint_changes_with_inputs(self):
        config = {'cx_b': {'prompt_method': 'A3-2'}}
        first = config_fingerprint(config, {'seed': 1, 'split': 'a'})
        second = config_fingerprint(config, {'seed': 2, 'split': 'a'})
        third = config_fingerprint(config, {'seed': 1, 'split': 'b'})
        self.assertNotEqual(first, second)
        self.assertNotEqual(first, third)

    def test_synthetic_short_optimization_changes_only_module(self):
        module = A3PromptAdapter(.5, 1., 0., trainable=True)
        unrelated = torch.nn.Parameter(torch.tensor(3.))
        heads = dict(semantic_score_map=torch.full((3, 3), .8),
                     instance_score_map=torch.full((3, 3), .2),
                     presence_score=torch.tensor(.8))
        target = torch.ones(3, 3)
        valid = torch.ones(3, 3, dtype=torch.bool)
        def samples():
            yield module(heads), target, valid
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            result = optimize_binary_module(
                module, samples, samples, IndependentBinaryLoss('A3-2'),
                dict(lr=.01, epochs=2, patience=0,
                     validate_every_epochs=1, param_reg_weight=0), .5,
                log_prefix='prompt test')
        self.assertEqual(float(unrelated), 3.)
        self.assertEqual(result['status'], 'trained')
        self.assertIsNotNone(result['initial_metrics'])
        self.assertAlmostEqual(
            result['iou_gain_vs_initial'],
            result['best_metrics']['IoU'] - result['initial_metrics']['IoU'])
        for metric in ('IoU', 'Acc', 'Precision', 'Recall', 'Fscore'):
            self.assertIn(metric, result['metric_gains'])
        self.assertIn('[prompt test] validation epoch=0', output.getvalue())
        self.assertIn('[prompt test] selected epoch=', output.getvalue())
        self.assertIn('expected_finish=', output.getvalue())

    def test_train_pipeline_saves_periodic_and_last_loadable_pth(self):
        module = A3PromptAdapter(.5, 1., 0., trainable=True)
        heads = dict(
            semantic_score_map=torch.full((3, 3), .8),
            instance_score_map=torch.full((3, 3), .2),
            presence_score=torch.tensor(.8))
        target = torch.ones(3, 3)
        valid = torch.ones(3, 3, dtype=torch.bool)

        def samples():
            yield module(heads), target, valid

        with tempfile.TemporaryDirectory() as directory:
            trainer = CXBTrainer(
                {'cx_b': {'prompt_init': 'scratch'}}, __file__, directory)
            checkpoint_dir = Path(directory) / 'prompt_000'
            callback = trainer._training_checkpoint_callback(
                checkpoint_dir,
                dict(kind='prompt', method='A3-2',
                     base_parameters=DEFAULT_A3_PARAMETERS,
                     base_trainable=False, hidden_channels=16,
                     identity={'prompt': 'building'}))
            optimize_binary_module(
                module, samples, samples, IndependentBinaryLoss('A3-2'),
                dict(lr=.01, epochs=3, patience=0,
                     validate_every_epochs=1, checkpoint_interval_epochs=2,
                     save_last_checkpoint=True, param_reg_weight=0),
                checkpoint_callback=callback)
            periodic = checkpoint_dir / 'epoch_0002.pth'
            last = checkpoint_dir / 'last.pth'
            self.assertTrue(periodic.is_file())
            self.assertTrue(last.is_file())
            _, periodic_metadata = load_variant(periodic)
            _, last_metadata = load_variant(last)
            self.assertEqual(periodic_metadata['epoch'], 2)
            self.assertEqual(periodic_metadata['checkpoint_type'], 'periodic')
            self.assertEqual(last_metadata['epoch'], 3)
            self.assertEqual(last_metadata['checkpoint_type'], 'last')
            payload = torch.load(last, map_location='cpu')
            self.assertEqual(payload['training_state']['epoch'], 3)
            self.assertIn(
                'optimizer_state_dict', payload['training_state'])

    def test_all_combinations_have_a_synthetic_training_path(self):
        for prompt_method, concept_method, loss_method in COMBINATION_CONFIGS:
            prompt_modules = []
            probabilities = []
            semantics, instances, presences = [], [], []
            regions = []
            for offset in (0., .1):
                heads = dict(
                    semantic_score_map=torch.full((8, 9), .7-offset),
                    instance_score_map=torch.full((8, 9), .3+offset),
                    presence_score=torch.tensor(.8-offset))
                module = build_prompt_module(
                    prompt_method, DEFAULT_A3_PARAMETERS, trainable=True)
                prompt_modules.append(module)
                probabilities.append(module(heads))
                semantics.append(heads['semantic_score_map'])
                instances.append(heads['instance_score_map'])
                presences.append(heads['presence_score'])
                regions.append((
                    torch.ones(1, 8, 9, dtype=torch.bool),
                    torch.tensor([.8])))
            probabilities = torch.stack(probabilities)
            aggregator = build_concept_aggregator(concept_method, 2)
            output = aggregator(
                probabilities,
                presence=torch.stack(presences),
                heads=dict(semantic_score_map=torch.stack(semantics),
                           instance_score_map=torch.stack(instances)),
                regions=regions)
            config, _ = resolve_combination_config(prompt_method, concept_method, loss_method)
            loss = build_binary_loss(config, "concept")(
                output, torch.ones_like(output),
                torch.ones_like(output, dtype=torch.bool))
            trainable = [parameter for module in [*prompt_modules, aggregator]
                         for parameter in module.parameters()
                         if parameter.requires_grad]
            if trainable:
                loss.backward()
            self.assertTrue(torch.isfinite(loss), (
                prompt_method, concept_method, loss_method))


if __name__ == '__main__':
    unittest.main()
