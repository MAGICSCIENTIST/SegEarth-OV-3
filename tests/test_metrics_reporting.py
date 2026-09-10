from pathlib import Path

import torch
from openpyxl import load_workbook

from custom_metrics import PromptIoUMetric
from metrics_reporting import write_metrics_workbook


def test_prompt_metric_uses_zero_based_concept_index_for_plain_txt(tmp_path):
    concepts = tmp_path / 'concepts.txt'
    concepts.write_text('background\nbuilding,house\n', encoding='utf-8')
    metric = PromptIoUMetric(
        str(concepts), score_threshold=0.5, ignore_index=255)
    # GT: background, building, ignored, building.
    gt = torch.tensor([[[0, 1], [255, 1]]])
    scores = torch.tensor([
        [[0.9, 0.2], [0.9, 0.1]],  # background: TP=1, no FP
        [[0.1, 0.9], [0.9, 0.2]],  # building: TP=1, FN=1
        [[0.1, 0.8], [0.9, 0.7]],  # house: TP=2
    ])
    metric.process({}, [{
        'gt_sem_seg': {'data': gt},
        'raw_prompt_scores': {'data': scores},
    }])
    assert metric.compute_metrics(metric.results) == {}
    rows = {row['Prompt']: row for row in metric._last_prompt_metrics}
    assert rows['background']['IoU'] == 100.0
    assert rows['building']['TP'] == 1
    assert rows['building']['FN'] == 1
    assert rows['building']['Recall'] == 50.0
    assert rows['house']['IoU'] == 100.0


def test_workbook_preserves_legacy_columns_and_adds_prompt_sheet(tmp_path):
    output = tmp_path / 'results.xlsx'
    write_metrics_workbook(
        str(output),
        summaries=[{'mIoU': 47.41}],
        model='SAM3', dataset='LoveDA')
    write_metrics_workbook(
        str(output),
        prompt_metrics=[{
            'Prompt': 'building', 'ConceptID': 'building',
            'ConceptName': 'building', 'IoU': 61.5, 'TP': 10, 'FP': 2,
            'FN': 3, 'Threshold': 0.5,
        }],
        model='SAM3', dataset='LoveDA')
    workbook = load_workbook(output, read_only=True, data_only=True)
    assert 'summary' in workbook.sheetnames
    assert 'per_prompt' in workbook.sheetnames
    summary_headers = next(workbook['summary'].values)
    prompt_headers = next(workbook['per_prompt'].values)
    assert summary_headers[:8] == (
        'Model', 'Dataset', 'aAcc', 'mIoU', 'mAcc', 'mFscore',
        'mPrecision', 'mRecall')
    assert 'TP' in prompt_headers
    assert 'Threshold' in prompt_headers
