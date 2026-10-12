"""Validate the fixed report contract and its effect on training/evaluation."""
import contextlib
import csv
import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import torch
from PIL import Image
from openpyxl import load_workbook

from cx_b.configuration import resolve_combination_config
from cx_b.joint import JointLoss, optimize_joint
from cx_b.losses import IndependentBinaryLoss
from cx_b.modules import A3PromptAdapter
from cx_b.reporting import CXBReporter, LayerCounts, JointObserver, binary_metrics
from cx_b.report_workbook import write_best_workbook
from cx_b.trainer import CXBTrainer, optimize_binary_module
import eval_batch_cx_b as batch
import test_cx_b_joint as joint_fixtures
import test_eval_batch_cx_b as fixtures


def csv_groups(text):
    """Parse each printed CSV block exactly as copied from the TXT."""
    result, name, lines = [], None, []
    for line in text.splitlines():
        if line.startswith('['):
            if name:
                result.append((name, list(csv.DictReader(io.StringIO('\n'.join(lines))))))
            name = next((k for k in ('prompt', 'concept', 'summary') if line.startswith('['+k+']')), None)
            lines = []
        elif name:
            lines.append(line)
    if name:
        result.append((name, list(csv.DictReader(io.StringIO('\n'.join(lines))))))
    return result


def run_training_fixture(root, mode='independent'):
    """Run the real orchestrator on cached, tiny data; keep SAM construction mocked."""
    directory = fixtures.BatchEvaluationTests().fixture(root)
    snapshot = batch.read_json(directory/'run_config.json')
    config = snapshot['config']
    common, _ = resolve_combination_config('A3-2', 'dynamic', 'A3-2')
    config['cx_b'] = common['cx_b']
    config['cx_b'].update(joint_mode=mode, prediction_background=0)
    config['model']['prob_thd'] = .5
    for stage in ('prompt_train', 'concept_train', 'joint_train'):
        config['cx_b'][stage].update(epochs=1, validate_every_epochs=1, checkpoint_interval_epochs=1)
    source = root/'concepts.txt'
    source.write_text(snapshot['inputs']['concept_content'], encoding='utf-8')
    raw = torch.tensor(np.asarray(Image.open(root/'dataset/gt/sample.png')).copy()).long()
    sample = dict(data_samples=SimpleNamespace(metainfo=dict(img_path=str(root/'dataset/img/sample.png')),
                                              gt_sem_seg=SimpleNamespace(data=raw.unsqueeze(0))))
    heads = {i: torch.load(next((root/'cache/val'/f'prompt_{i:03d}').glob('*.pth')))['heads'] for i in range(3)}
    provider = SimpleNamespace(load=lambda split, prompt, index, image: heads[prompt])
    model = SimpleNamespace(device=torch.device('cpu'))
    trainer = CXBTrainer(config, source, directory, resume=False,
                         cli_args=dict(joint_mode=mode, seed=3407, description='synthetic report check'))
    console = io.StringIO()
    with contextlib.redirect_stdout(console), \
         patch.object(trainer, '_resolved_paths', return_value=(str(source), str(root/'cache'))), \
         patch.object(trainer, 'run_fingerprint', return_value=('report-fixture', snapshot['inputs'])), \
         patch.object(trainer, '_build_runtime', return_value=(model, [sample], [sample], provider)):
        result = trainer.run()
    return directory, config, console.getvalue(), result


class ReportingTests(unittest.TestCase):
    def reporter(self, root):
        report = CXBReporter(root, dict(model=dict(prob_thd=.5), cx_b=dict(metric_threshold=.5,
            prediction_background=0, joint_mode='independent')), dict(seed=3407, items=['a', 'b']))
        report.catalog([dict(concept_id='building'), dict(concept_id='road')], [
            dict(prompt_index=0, prompt='building, "house"', concept_index=0, concept_id='building'),
            dict(prompt_index=1, prompt='road', concept_index=1, concept_id='road')])
        return report

    def test_json_csv_units_nulls_and_file_roles(self):
        with tempfile.TemporaryDirectory() as temp:
            root, console = Path(temp), io.StringIO()
            report = self.reporter(root)
            with contextlib.redirect_stdout(console), report:
                report.phase('prompt', 'building')
                report.phase('prompt', 'building')
                report.progress_line('train', 1, 2, report.tables({0: binary_metrics(80, 20, 40)}), .25, 1, 2)
                report.validation(1, 2, report.tables({0: binary_metrics(80, 20, 40)}))
                report.start()
            params = (root/'log_fix.txt').read_text(encoding='utf-8')
            val = (root/'val_fix.txt').read_text(encoding='utf-8')
            args = json.loads(params.split('[args]', 1)[1].split('\n', 1)[1].split('[config]')[0])
            config = json.loads(params.split('[config]', 1)[1].split('\n', 1)[1].split('[metric_definition]')[0])
            self.assertEqual(args, report.args)
            self.assertEqual(config, report.config)
            self.assertEqual(params.count('[args]'), 1)
            self.assertNotIn('[epoch=', params)
            self.assertNotIn('[train', val)
            self.assertNotIn('validation_id', console.getvalue())
            self.assertEqual(val.count('[stage=prompt'), 1)
            groups = csv_groups(val)
            self.assertEqual([name for name, _ in groups], ['prompt', 'concept', 'summary'])
            prompt, concept, summary = [rows for _, rows in groups]
            self.assertEqual(prompt[0]['prompt'], 'building, "house"')
            self.assertEqual(prompt[0]['iou'], '57.1429')
            self.assertEqual(prompt[0]['precision'], '80.0000')
            self.assertEqual(prompt[0]['TP'], '80')
            self.assertEqual(prompt[1]['TP'], 'null')
            self.assertTrue(all(row['iou']=='null' for row in concept))
            self.assertEqual([row['level'] for row in summary], ['multiclass', 'concept_mean', 'prompt_mean'])
            self.assertEqual(summary[0]['iou'], 'null')
            self.assertEqual(summary[2]['TP'], '80')
            for row in prompt+concept+summary:
                self.assertNotIn('iou_prompt', row)
                self.assertNotIn('iou_concept', row)
                self.assertRegex(row['elapsed'], r'^\d{2,}:\d{2}:\d{2}$')
                self.assertIn('expected_finish', row)
            self.assertEqual({p.name for p in root.iterdir()}, {'log_fix.txt', 'val_fix.txt'})

    def test_timing_covers_dependency_stdout_and_stderr_without_corrupting_json(self):
        import sys
        with tempfile.TemporaryDirectory() as temp:
            stdout, stderr = io.StringIO(), io.StringIO()
            with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
                with self.reporter(Path(temp)) as report:
                    report.begin_run()
                    print('dependency elapsed=00:01 eta=00:02 expected_finish=2020-01-01 01:02:03')
                    print('diagnostic', file=sys.stderr)
                self.assertIs(sys.stdout, stdout)
                self.assertIs(sys.stderr, stderr)
            line = stdout.getvalue().splitlines()[-1]
            self.assertEqual(line.count('elapsed='), 1)
            self.assertNotIn('2020-01-01', line)
            self.assertIn('expected_finish=', stderr.getvalue())

    def test_summary_uses_macro_ratios_summed_counts_and_individual_prompts(self):
        with tempfile.TemporaryDirectory() as temp:
            report = self.reporter(Path(temp))
            report.records.append(dict(prompt_index=2, prompt='house',
                                       concept_index=0, concept_id='building'))
            tables = report.tables(
                {0: binary_metrics(1, 0, 0), 1: binary_metrics(0, 100, 1),
                 2: binary_metrics(0, 0, 1000)},
                {0: binary_metrics(0, 0, 10)})
            multiclass, concepts, prompts = tables[2]
            self.assertEqual([row['level'] for row in tables[2]],
                             ['multiclass', 'concept_mean', 'prompt_mean'])
            self.assertTrue(all(value is None for key, value in multiclass.items() if key != 'level'))
            self.assertEqual(concepts['iou'], 0)
            self.assertIsNone(concepts['precision'])
            self.assertEqual(concepts['recall'], 0)
            self.assertEqual(concepts['f1'], 0)
            self.assertEqual(concepts['FN'], 10)
            # Three prompts: (100 + 0 + 0)/3. Averaging by concept would give 25.
            self.assertAlmostEqual(prompts['iou'], 100/3)
            self.assertAlmostEqual(prompts['recall'], 100/3)
            self.assertAlmostEqual(prompts['f1'], 100/3)
            self.assertEqual(prompts['precision'], 50)  # Undefined precision excluded.
            self.assertEqual([prompts[k] for k in ('TP', 'FP', 'FN', 'PredPixels', 'GTPixels')],
                             [1, 100, 1001, 101, 1002])
            self.assertNotAlmostEqual(prompts['iou'], 100/1102)  # Not micro IoU.

    def test_each_report_rebuilds_required_summary_from_current_details(self):
        with tempfile.TemporaryDirectory() as temp:
            root, console = Path(temp), io.StringIO()
            report = self.reporter(root)
            with contextlib.redirect_stdout(console), report:
                report.phase('joint')
                previous = report.tables({0: binary_metrics(1, 0, 0)})
                # Missing summary must not omit the third CSV table.
                report.validation(1, 2, previous[:2])
                current = report.tables({0: binary_metrics(1, 1, 0)},
                                        {0: binary_metrics(1, 0, 1)},
                                        {0: binary_metrics(1, 1, 1)})
                # A stale best/previous-epoch summary must not leak into final output.
                report.validation(0, 0, (*current[:2], previous[2]), final=True, best_epoch=2)
            saved = (root/'val_fix.txt').read_text(encoding='utf-8')
            self.assertEqual(csv_groups(console.getvalue()), csv_groups(saved))
            groups = csv_groups(saved)
            self.assertEqual([name for name, _ in groups], ['prompt', 'concept', 'summary']*2)
            self.assertEqual(groups[2][1][2]['iou'], '100.0000')
            final = groups[-1][1]
            self.assertEqual([row['level'] for row in final],
                             ['multiclass', 'concept_mean', 'prompt_mean'])
            self.assertEqual([row['iou'] for row in final], ['33.3333', '50.0000', '50.0000'])
            self.assertIn('[final best_epoch=2]', saved)
            self.assertEqual(list(final[0]), ['level', 'iou', 'precision', 'recall', 'f1',
                'TP', 'FP', 'FN', 'PredPixels', 'GTPixels', 'elapsed', 'eta', 'expected_finish'])

    def test_skip_preserves_reports_and_real_restart_resets_only_fixed_reports(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for name in ('log_fix.txt', 'val_fix.txt', 'results.txt'):
                (root/name).write_text('previous '+name)
            with contextlib.redirect_stdout(io.StringIO()), self.reporter(root):
                print('Matching completed run found; training skipped.')
            self.assertEqual((root/'log_fix.txt').read_text(), 'previous log_fix.txt')
            self.assertEqual((root/'val_fix.txt').read_text(), 'previous val_fix.txt')
            with contextlib.redirect_stdout(io.StringIO()), self.reporter(root) as report:
                report.phase('prompt', 'building')
                report.validation(0, 1, report.tables())
            self.assertNotIn('previous', (root/'val_fix.txt').read_text())
            self.assertEqual((root/'results.txt').read_text(), 'previous results.txt')

    def test_all_layers_match_batch_eval_with_ignore_and_unassigned_pixels(self):
        records = [dict(concept_index=0), dict(concept_index=1)]
        scores = torch.tensor([[[.1, .6, .1, .99]], [[.2, .2, .9, .01]]])
        target = torch.tensor([[0, 0, 1, -1]])
        valid = target >= 0
        counters = LayerCounts(records, 2, .5, .5, 65535)
        counters.update(scores, list(scores), target, valid)
        result = counters.result()
        confusion = np.zeros((2, 3), np.int64)
        prediction = batch.classify_scores(scores.numpy(), [0, 1], .5, batch.NODATA)
        batch.add_confusion(confusion, prediction, target.numpy(), valid.numpy())
        reference = batch.summarize_confusion(confusion, ['bg', 'fg'])
        self.assertEqual(result['confusion_matrix'], confusion.tolist())
        self.assertEqual(result['unassigned_pixels'], 1)
        for key, value in reference['summary'].items():
            self.assertAlmostEqual(result[key]*100, value)
        self.assertEqual(result['prompt_binary'], result['concept_binary'])
        self.assertEqual(result['per_class'][0]['FN'], 1)

    def test_observation_does_not_change_binary_training_or_selection(self):
        with tempfile.TemporaryDirectory() as temp:
            states, results = [], []
            for observed in (False, True):
                module = A3PromptAdapter(trainable=True)
                heads = dict(semantic_score_map=torch.full((3, 3), .2),
                             instance_score_map=torch.full((3, 3), .4), presence_score=torch.tensor(.8))
                def samples():
                    yield module(heads), torch.ones(3, 3), torch.ones(3, 3, dtype=torch.bool)
                report = self.reporter(Path(temp))
                with contextlib.redirect_stdout(io.StringIO()):
                    observer = report.binary('prompt', 0, 3, 1) if observed else None
                    results.append(optimize_binary_module(module, samples, samples, IndependentBinaryLoss(),
                        dict(epochs=3, lr=.1, patience=0), observer=observer))
                states.append(module.state_dict())
            # Legacy binary metrics can contain NaN for undefined precision.
            # JSON preserves their positions while comparing every other value exactly.
            self.assertEqual(json.dumps(results[0], sort_keys=True), json.dumps(results[1], sort_keys=True))
            for name in states[0]:
                self.assertTrue(torch.equal(states[0][name], states[1][name]))

    def test_observation_does_not_change_joint_training_or_selection(self):
        with tempfile.TemporaryDirectory() as temp:
            states, results = [], []
            for observed in (False, True):
                pipeline, heads, target = joint_fixtures.JointTests().pipeline('multiclass_joint')
                samples = lambda: iter([(heads, None, target, target >= 0)])
                config = dict(epochs=2, lr=.01, patience=0)
                report = self.reporter(Path(temp))
                records = [dict(r, prompt=str(r['prompt_index']), concept_id=str(r['concept_index'])) for r in pipeline.records]
                report.catalog([dict(concept_id=str(i)) for i in range(2)], records)
                with contextlib.redirect_stdout(io.StringIO()):
                    observer = JointObserver(report, 2, 1) if observed else None
                    results.append(optimize_joint(pipeline, samples, samples,
                        JointLoss(IndependentBinaryLoss(), pipeline.mode, config), config, observer=observer))
                states.append(pipeline.state_dict())
            self.assertEqual(results[0], results[1])
            for name in states[0]:
                self.assertTrue(torch.equal(states[0][name], states[1][name]))

    def test_independent_and_joint_runs_only_write_new_text_reports_and_match_final_eval(self):
        for mode in ('independent', 'multiclass_joint'):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as temp:
                directory, _, console, result = run_training_fixture(Path(temp), mode)
                report = (directory/'val_fix.txt').read_text(encoding='utf-8')
                self.assertEqual([p.name for p in directory.rglob('*.xlsx')], ['result_fix.xlsx'])
                self.assertEqual(len(list(directory.rglob('*.log'))), 0)
                self.assertEqual({p.name for p in directory.rglob('*.txt')}, {'log_fix.txt', 'val_fix.txt'})
                for line in console.splitlines():
                    if line.startswith('['):
                        self.assertIn('elapsed=', line)
                        self.assertIn('eta=', line)
                        self.assertIn('expected_finish=', line)
                marker = '[final best_epoch=per_module]' if mode=='independent' else '[final best_epoch=0]'
                self.assertIn(marker, report)
                groups = csv_groups(report)
                self.assertEqual(len(groups)%3, 0)
                for _, rows in groups:
                    self.assertTrue(rows)
                    self.assertTrue(all(row.get('iou') is not None for row in rows))
                final = groups[-1][1]
                args = batch.parse_args(['--root', str(directory), '--device=cpu', '--out-root', str(Path(temp)/'eval')])
                with contextlib.redirect_stdout(io.StringIO()):
                    evaluated = batch.evaluate_run(directory, args)['result']
                self.assertAlmostEqual(float(final[0]['iou']), evaluated['metrics']['multiclass_saved']['summary']['mIoU'], places=4)
                self.assertAlmostEqual(float(final[1]['iou']), evaluated['metrics']['independent_binary']['summary']['mIoU'], places=4)
                self.assertAlmostEqual(float(final[2]['iou']), evaluated['metrics']['prompt_binary_ge']['summary']['mIoU'], places=4)
                self.assertEqual(result['summary']['iou'], 100)
                original_txt = {name: (directory/name).read_bytes() for name in ('log_fix.txt', 'val_fix.txt')}
                # Rebuilding is best-only, never appends validation history or edits TXT.
                write_best_workbook(directory)
                for name, content in original_txt.items():
                    self.assertEqual((directory/name).read_bytes(), content)
                workbook = load_workbook(directory/'result_fix.xlsx', read_only=True, data_only=True)
                try:
                    self.assertEqual(workbook.sheetnames, ['summary', 'per_class', 'per_prompt'])
                    summary_rows = list(workbook['summary'].values)
                    self.assertEqual(summary_rows[0][:4], ('Model', 'Dataset', 'level', 'iou'))
                    self.assertEqual(len(summary_rows), 4)
                    for row, text_row in zip(summary_rows[1:], final):
                        mapped = dict(zip(summary_rows[0], row))
                        self.assertEqual(mapped['level'], text_row['level'])
                        for metric in ('iou', 'precision', 'recall', 'f1'):
                            self.assertAlmostEqual(mapped[metric], float(text_row[metric]), places=4)
                        for metric in ('TP', 'FP', 'FN', 'PredPixels', 'GTPixels'):
                            self.assertEqual(mapped[metric], int(text_row[metric]))
                    self.assertEqual(workbook['per_class'].max_row, 5)
                    self.assertEqual(workbook['per_prompt'].max_row, 4)
                    for name in ('per_class', 'per_prompt'):
                        rows = list(workbook[name].values)
                        self.assertEqual(rows[0][:2], ('Model', 'Dataset'))
                        self.assertIn('BestEpoch', rows[0])
                        self.assertIn('InitialIoU', rows[0])
                finally:
                    workbook.close()


if __name__ == '__main__':
    unittest.main()
