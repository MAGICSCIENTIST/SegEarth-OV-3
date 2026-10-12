"""Export the restored best model to one workbook, without changing TXT reports.

Can also rebuild the workbook from saved final metrics without SAM or training:
python -m cx_b.report_workbook --run-dir PATH [--output PATH]
"""
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import re
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font

from metrics_reporting import CLASS_HEADERS, PROMPT_HEADERS
from .reporting import CXBReporter, FIELDS, TIMES, normalize, summary_rows


def read_json(path, default=None):
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else default


def indexed(value):
    if isinstance(value, dict):
        return {int(k): v for k, v in value.items()}
    return dict(enumerate(value or []))


def write_best_workbook(run_dir, output=None):
    """Overwrite a single best-only workbook using final_metrics.json as truth."""
    root = Path(run_dir)
    saved = read_json(root/'final_metrics.json')
    snapshot = read_json(root/'run_config.json')
    if not saved or not snapshot:
        raise ValueError('Best-model evaluation is required: final_metrics.json and run_config.json')
    if saved.get('scale') != 'fraction':
        raise ValueError('Expected fraction-scale final_metrics.json')
    config, metrics, best_epoch = snapshot['config'], saved['metrics'], saved['best_epoch']
    prompt_meta = sorted(read_json(root/'prompt_metrics.json', []), key=lambda r: r['prompt_index'])
    concept_meta = {r['concept_index']: r for r in read_json(root/'concept_metrics.json', [])}
    concepts = config.get('concepts') or [dict(concept_id=name, concept_name=name)
        for name in dict.fromkeys(r['concept_id'] for r in prompt_meta)]
    if not concepts or not prompt_meta:
        raise ValueError('Missing prompt/concept catalog for final metrics')
    concept_indices = {c['concept_id']: i for i, c in enumerate(concepts)}
    records = [dict(r, concept_index=concept_indices[r['concept_id']]) for r in prompt_meta]
    report = CXBReporter(root, config)
    report.catalog(concepts, records)
    tables = report.tables(indexed(metrics.get('prompt_binary')), indexed(metrics.get('concept_binary')),
                           indexed(metrics.get('per_class')))
    cx = config['cx_b']
    common = dict(Model=f'CX-B:{cx["prompt_method"]}+{cx["concept_method"]}',
                  Dataset=config.get('dataset_type') or config.get('validation_dataloader', {}).get(
                      'dataset', {}).get('type', 'unknown'))
    times = dict.fromkeys(TIMES)
    txt = root/'val_fix.txt'
    if txt.exists():
        final_lines = [line for line in txt.read_text(encoding='utf-8').splitlines()
                       if line.startswith('[final best_epoch=')]
        if final_lines:
            for key in TIMES:
                pattern = (r'expected_finish=(pending|\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})'
                           if key == 'expected_finish' else key+r'=(\S+)')
                match = re.search(pattern, final_lines[-1])
                times[key] = match.group(1) if match else None
    joint_initial = (read_json(root/'joint'/'training_metrics.json', {}).get('initial_metrics', {})
                     if cx.get('joint_mode', 'independent') != 'independent' else {})

    def detail(row, epoch, threshold, initial=None, maximum=None):
        aliases = dict(IoU='iou', Acc='recall', Precision='precision', Recall='recall', Fscore='f1')
        out = dict(common, **{k: row[v] for k, v in aliases.items()},
                   **{k: row[k] for k in ('TP', 'FP', 'FN', 'PredPixels', 'GTPixels')},
                   Epoch=epoch, BestEpoch=epoch, Split='validation', Threshold=threshold)
        out['PredGTAreaRatio'] = (row['PredPixels']/row['GTPixels']
                                  if row['GTPixels'] and row['PredPixels'] is not None else None)
        if initial:
            normalized = normalize(initial)
            for key, normalized_key in aliases.items():
                baseline = normalized[normalized_key]
                out['Initial'+key] = baseline
                out[key+'Gain'] = (out[key]-baseline if out[key] is not None and baseline is not None else None)
            out['InitialTP'] = initial.get('TP')
        if maximum:
            out['MaxBaselineIoU'] = normalize(maximum)['iou']
            out['IoUDeltaVsMax'] = (row['iou']-out['MaxBaselineIoU']
                if row['iou'] is not None and out['MaxBaselineIoU'] is not None else None)
        return out

    class_rows = []
    for row in tables[1]:
        c = concept_indices[row['concept']]
        meta = concept_meta.get(c, {})
        level = row['level']
        initial, maximum = None, None
        if joint_initial:
            source = 'per_class' if level == 'multiclass' else 'concept_binary'
            initial = indexed(joint_initial.get(source)).get(c)
        elif level == 'concept' and meta.get('metric_scope') != 'final_multiclass':
            initial, maximum = meta.get('initial_metrics'), meta.get('max_baseline')
        epoch = (meta.get('best_epoch') if best_epoch == 'per_module' and level == 'concept' else best_epoch)
        threshold = cx.get('metric_threshold', .5) if level == 'concept' else config['model'].get('prob_thd', .1)
        class_rows.append(dict(detail(row, epoch, threshold, initial, maximum),
                               Class=concepts[c].get('concept_name') or row['concept'], level=level))
    prompt_rows = []
    for row, meta in zip(tables[0], prompt_meta):
        initial = (indexed(joint_initial.get('prompt_binary')).get(meta['prompt_index'])
                   if joint_initial else meta.get('initial_metrics'))
        epoch = meta.get('best_epoch') if best_epoch == 'per_module' else best_epoch
        c = concepts[concept_indices[row['concept']]]
        prompt_rows.append(dict(detail(row, epoch, cx.get('metric_threshold', .5), initial),
            Prompt=row['prompt'], ConceptID=row['concept'], ConceptName=c.get('concept_name') or row['concept']))

    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = dict(summary=['Model', 'Dataset', 'level', *FIELDS, *TIMES],
                   per_class=[*CLASS_HEADERS[:3], 'level', *CLASS_HEADERS[3:]],
                   per_prompt=PROMPT_HEADERS)
    all_rows = dict(summary=[dict(common, **row, **times) for row in summary_rows(*tables[:2])],
                    per_class=class_rows, per_prompt=prompt_rows)
    for name, columns in headers.items():
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in all_rows[name]:
            values = [row.get(k) for k in columns]
            sheet.append(['null' if v is None or isinstance(v, float) and not math.isfinite(v) else v
                          for v in values])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            sheet.column_dimensions[cell.column_letter].width = max(16, len(str(cell.value))+2)
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, float):
                    cell.number_format = '0.0000'
    destination = Path(output) if output else root/'result_fix.xlsx'
    destination.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(suffix='.xlsx', prefix='.result_fix_', dir=destination.parent)
    os.close(handle)
    try:
        workbook.save(temporary)
        os.replace(temporary, destination)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir', required=True, type=Path)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    print(write_best_workbook(args.run_dir, args.output))


if __name__ == '__main__':
    main()
